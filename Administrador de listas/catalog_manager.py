from __future__ import annotations

import copy
import json
import re
import shutil
import tempfile
import unicodedata
import uuid
from datetime import datetime
from pathlib import Path
from urllib.parse import parse_qs, urlparse

try:
    from openpyxl import Workbook, load_workbook
except ImportError:  # pragma: no cover - handled at runtime by CatalogError.
    Workbook = None
    load_workbook = None

CUSTOM_LISTS_KEY = "listasPersonalizadas"
COLLECTION_ORDER_KEY = "ordenColecciones"
EXPLANATORY_GUIDES_KEY = "guiasVisuales"
MEDIA_TYPES = {"imagen", "video"}
YOUTUBE_HOSTS = {"youtube.com", "www.youtube.com", "m.youtube.com", "youtu.be", "www.youtu.be"}


class CatalogError(ValueError):
    """Raised when the catalog does not meet the expected structure."""


class CatalogManager:
    def __init__(self, project_root: str | Path):
        selected_path = Path(project_root).expanduser().resolve()
        if selected_path.is_file():
            self.json_path = selected_path
            self.project_root = selected_path.parent.parent.parent
        else:
            selected_json = selected_path / "productos.json"
            if selected_json.is_file():
                self.json_path = selected_json
                self.project_root = selected_path.parent.parent
            else:
                self.project_root = selected_path
                self.json_path = self.project_root / "src" / "data" / "productos.json"
        self.images_root = self.project_root / "public" / "images"
        self.backups_root = self.project_root / "Administrador de listas" / "backups"
        self.originals_root = self.project_root / "Administrador de listas" / "originales"
        self.last_image_optimization: dict[str, int | str] | None = None
        self.data: dict | None = None

    def load(self) -> dict:
        try:
            with self.json_path.open("r", encoding="utf-8") as source:
                data = json.load(source)
        except FileNotFoundError as error:
            raise CatalogError(f"No se encontro el catalogo: {self.json_path}") from error
        except json.JSONDecodeError as error:
            raise CatalogError(f"El JSON no es valido: {error}") from error

        self.ensure_uids(data)
        self.validate(data)
        self.data = data
        return data

    def export_prices_excel(self, output_path: str | Path) -> str:
        if Workbook is None:
            raise CatalogError("Falta openpyxl. Instala las dependencias con: python -m pip install -r requirements.txt")
        if self.data is None:
            self.data = self.load()

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Precios"
        sheet.append(["coleccion", "id", "nombre", "precio"])

        for collection_name, items in self.data.items():
            if collection_name in {COLLECTION_ORDER_KEY, CUSTOM_LISTS_KEY, EXPLANATORY_GUIDES_KEY}:
                continue
            if not isinstance(items, list):
                continue
            for item in items:
                if not isinstance(item, dict):
                    continue
                if "id" not in item or "nombre" not in item:
                    continue
                price = item.get("precio")
                sheet.append([collection_name, item.get("id"), item.get("nombre"), "" if price is None else price])

        target = Path(output_path).expanduser().resolve()
        target.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(target)
        return str(target)

    def import_prices_excel(self, input_path: str | Path) -> dict[str, int]:
        if load_workbook is None:
            raise CatalogError("Falta openpyxl. Instala las dependencias con: python -m pip install -r requirements.txt")
        if self.data is None:
            self.data = self.load()

        workbook = load_workbook(str(Path(input_path).expanduser().resolve()), read_only=False, data_only=True)
        if not workbook.worksheets:
            raise CatalogError("El archivo Excel no contiene hojas válidas")
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            raise CatalogError("El archivo Excel está vacío")

        header_map = {}
        for index, cell in enumerate(rows[0]):
            label = str(cell or "").strip().lower()
            if not label:
                continue
            normalized = label.replace(" ", "").replace("-", "").replace("_", "")
            header_map[normalized] = index

        required = {"coleccion": "coleccion", "lista": "coleccion", "categoria": "coleccion", "id": "id", "producto": "nombre", "nombre": "nombre", "precio": "precio"}
        if not set(required).intersection(header_map):
            raise CatalogError("La hoja no tiene columnas válidas. Usa: coleccion, id, nombre, precio.")

        updated = 0
        skipped = 0
        for row in rows[1:]:
            if not row or not any(cell not in (None, "") for cell in row):
                continue
            try:
                collection_name = row[header_map.get("coleccion", header_map.get("lista", header_map.get("categoria", -1)))]
            except IndexError:
                skipped += 1
                continue
            if collection_name is None or str(collection_name).strip() == "":
                skipped += 1
                continue

            raw_id = row[header_map.get("id", -1)] if "id" in header_map else None
            if raw_id in (None, ""):
                skipped += 1
                continue
            try:
                product_id = int(raw_id)
            except (TypeError, ValueError):
                skipped += 1
                continue

            collection = self.data.get(str(collection_name).strip())
            if not isinstance(collection, list):
                skipped += 1
                continue

            target = next((item for item in collection if isinstance(item, dict) and item.get("id") == product_id), None)
            if target is None:
                skipped += 1
                continue

            raw_price = row[header_map.get("precio", -1)] if "precio" in header_map else None
            if raw_price in (None, ""):
                continue
            if isinstance(raw_price, str) and raw_price.strip() == "----":
                target["precio"] = "----"
                updated += 1
                continue
            try:
                value = float(raw_price)
            except (TypeError, ValueError):
                raise CatalogError(f"El precio '{raw_price}' no es numérico para el producto ID {product_id}")
            target["precio"] = int(value) if value.is_integer() else value
            updated += 1

        workbook.close()
        return {"updated": updated, "skipped": skipped}

    @staticmethod
    def validate(data: object) -> None:
        if not isinstance(data, dict):
            raise CatalogError("El catalogo debe ser un objeto JSON")

        for collection_name, collection in data.items():
            if collection_name == COLLECTION_ORDER_KEY:
                if not isinstance(collection, list) or not all(isinstance(item, str) and item.strip() for item in collection):
                    raise CatalogError("'ordenColecciones' debe ser una lista de claves no vacías")
                continue
            if collection_name == EXPLANATORY_GUIDES_KEY:
                CatalogManager.validate_explanatory_guides(collection)
                continue
            if not isinstance(collection, list):
                raise CatalogError(f"La coleccion '{collection_name}' debe ser una lista")
            if collection_name == CUSTOM_LISTS_KEY:
                for custom_list in collection:
                    CatalogManager.validate_custom_list(custom_list)
            else:
                item_uids = set()
                for product in collection:
                    CatalogManager.validate_media_container(product)
                    CatalogManager.validate_optional_uid(product, item_uids)

    @staticmethod
    def validate_optional_uid(value: object, seen: set[str]) -> None:
        if not isinstance(value, dict) or "uid" not in value:
            return
        uid = value["uid"]
        if not isinstance(uid, str) or not re.fullmatch(r"[A-Za-z0-9_-]{16,64}", uid):
            raise CatalogError("Cada producto o fila debe tener un uid valido")
        if uid in seen:
            raise CatalogError(f"UID duplicado: {uid}")
        seen.add(uid)

    @staticmethod
    def new_uid() -> str:
        return uuid.uuid4().hex

    @staticmethod
    def ensure_uids(data: dict) -> None:
        for collection_name, collection in data.items():
            if collection_name in {COLLECTION_ORDER_KEY, EXPLANATORY_GUIDES_KEY} or not isinstance(collection, list):
                continue
            if collection_name == CUSTOM_LISTS_KEY:
                for custom_list in collection:
                    if not isinstance(custom_list, dict):
                        continue
                    for row in custom_list.get("filas", []):
                        if isinstance(row, dict) and not row.get("uid"):
                            row["uid"] = CatalogManager.new_uid()
                continue
            for product in collection:
                if isinstance(product, dict) and not product.get("uid"):
                    product["uid"] = CatalogManager.new_uid()

    @staticmethod
    def validate_media_container(value: object) -> None:
        if not isinstance(value, dict):
            raise CatalogError("Cada producto o fila debe ser un objeto")
        gallery = value.get("galeria", [])
        if not isinstance(gallery, list):
            raise CatalogError("El campo 'galeria' debe ser una lista")
        for media in gallery:
            CatalogManager.validate_media(media)

    @staticmethod
    def validate_media(media: object) -> None:
        if not isinstance(media, dict) or media.get("tipo") not in MEDIA_TYPES:
            raise CatalogError("Cada elemento multimedia debe tener tipo imagen o video")
        if not isinstance(media.get("src"), str) or not media["src"].strip():
            raise CatalogError("Cada elemento multimedia debe tener un src")

    @staticmethod
    def validate_explanatory_guides(value: object) -> None:
        if not isinstance(value, dict):
            raise CatalogError("'guiasVisuales' debe ser un objeto")
        for collection_key, guide in value.items():
            if not isinstance(collection_key, str) or not collection_key.strip():
                raise CatalogError("Cada guía visual debe tener una clave de lista válida")
            if not isinstance(guide, list):
                raise CatalogError(f"La guía visual '{collection_key}' debe ser una lista")
            for media in guide:
                if not isinstance(media, dict):
                    raise CatalogError("Cada elemento explicativo debe ser un objeto")
                if not isinstance(media.get("src"), str) or not media["src"].strip():
                    raise CatalogError("Cada elemento explicativo debe tener un src")
                if not isinstance(media.get("alt"), str) or not media["alt"].strip():
                    raise CatalogError("Cada elemento explicativo debe tener un alt")
                media_type = media.get("tipo", "imagen")
                if media_type not in MEDIA_TYPES:
                    raise CatalogError("Cada elemento explicativo debe tener tipo imagen o video")
                if media_type == "video":
                    CatalogManager.youtube_url(media["src"])

    @staticmethod
    def validate_custom_list(custom_list: object) -> None:
        if not isinstance(custom_list, dict):
            raise CatalogError("Cada lista personalizada debe ser un objeto")
        for field in ("clave", "titulo", "columnas", "filas"):
            if field not in custom_list:
                raise CatalogError(f"La lista personalizada no tiene '{field}'")
        if not isinstance(custom_list["clave"], str) or not custom_list["clave"].strip():
            raise CatalogError("La clave de una lista personalizada no puede estar vacia")
        if not isinstance(custom_list["columnas"], list) or not custom_list["columnas"]:
            raise CatalogError("Una lista personalizada debe tener columnas")
        column_keys = set()
        for column in custom_list["columnas"]:
            if not isinstance(column, dict) or not isinstance(column.get("clave"), str) or not isinstance(column.get("nombre"), str):
                raise CatalogError("Cada columna debe tener clave y nombre")
            if column["clave"] in column_keys:
                raise CatalogError(f"Columna duplicada: {column['clave']}")
            column_keys.add(column["clave"])
        if not isinstance(custom_list["filas"], list):
            raise CatalogError("Las filas deben ser una lista")
        row_uids = set()
        for row in custom_list["filas"]:
            if not isinstance(row, dict) or not isinstance(row.get("celdas"), dict):
                raise CatalogError("Cada fila debe tener un objeto celdas")
            if not set(row["celdas"]).issubset(column_keys):
                raise CatalogError("Una fila contiene una columna inexistente")
            CatalogManager.validate_media_container(row)
            CatalogManager.validate_optional_uid(row, row_uids)

    def save(self, data: dict) -> None:
        self.ensure_uids(data)
        self.validate(data)
        self.json_path.parent.mkdir(parents=True, exist_ok=True)
        if self.json_path.exists():
            self.backups_root.mkdir(parents=True, exist_ok=True)
            stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
            shutil.copy2(self.json_path, self.backups_root / f"productos-{stamp}.json")

        temporary_path: Path | None = None
        try:
            with tempfile.NamedTemporaryFile("w", encoding="utf-8", dir=self.json_path.parent, delete=False, suffix=".tmp") as target:
                json.dump(data, target, ensure_ascii=False, indent=2)
                target.write("\n")
                temporary_path = Path(target.name)
            temporary_path.replace(self.json_path)
        finally:
            if temporary_path and temporary_path.exists():
                temporary_path.unlink()

    def copy_image(self, source: str | Path, folder_key: str) -> str:
        source_path = Path(source).expanduser().resolve()
        if not source_path.is_file():
            raise CatalogError(f"No se encontro la imagen: {source_path}")
        try:
            from PIL import Image, ImageOps, UnidentifiedImageError
        except ImportError as error:
            raise CatalogError("Falta Pillow. Instala las dependencias con: python -m pip install -r requirements.txt") from error

        safe_folder = self.safe_key(folder_key)
        destination_folder = self.images_root / safe_folder
        destination_folder.mkdir(parents=True, exist_ok=True)
        originals_folder = self.originals_root / safe_folder
        originals_folder.mkdir(parents=True, exist_ok=True)
        original_destination = originals_folder / self.unique_filename(originals_folder, source_path.name)
        shutil.copy2(source_path, original_destination)

        destination_name = self.unique_filename(destination_folder, f"{source_path.stem}.webp")
        destination = destination_folder / destination_name
        try:
            with Image.open(source_path) as image:
                if getattr(image, "is_animated", False):
                    raise CatalogError("No se pueden optimizar GIF animados. Convierte la imagen a JPG o PNG antes de importarla.")
                image = ImageOps.exif_transpose(image)
                image.load()
                original_width, original_height = image.size
                image.thumbnail((1600, 1600), Image.Resampling.LANCZOS)
                if image.mode in ("RGBA", "LA") or "transparency" in image.info:
                    optimized = image.convert("RGBA")
                else:
                    optimized = image.convert("RGB")
                optimized.save(destination, "WEBP", quality=82, method=6)
        except UnidentifiedImageError as error:
            raise CatalogError(f"El archivo no es una imagen válida: {source_path}") from error
        except OSError as error:
            raise CatalogError(f"No se pudo optimizar la imagen: {error}") from error

        self.last_image_optimization = {
            "original_bytes": source_path.stat().st_size,
            "optimized_bytes": destination.stat().st_size,
            "original_width": original_width,
            "original_height": original_height,
            "optimized_width": image.width,
            "optimized_height": image.height,
            "original_name": source_path.name,
            "optimized_name": destination.name,
        }
        return f"/images/{safe_folder}/{destination_name}"

    def copy_explanatory_image(self, source: str | Path, collection_key: str) -> str:
        safe_collection = self.safe_key(collection_key)
        return self._copy_image_to_folders(
            source,
            self.images_root / "explicaciones" / safe_collection,
            self.originals_root / "explicaciones" / safe_collection,
            f"/images/explicaciones/{safe_collection}",
        )

    def _copy_image_to_folders(self, source: str | Path, destination_folder: Path, originals_folder: Path, public_prefix: str) -> str:
        source_path = Path(source).expanduser().resolve()
        if not source_path.is_file():
            raise CatalogError(f"No se encontro la imagen: {source_path}")
        try:
            from PIL import Image, ImageOps, UnidentifiedImageError
        except ImportError as error:
            raise CatalogError("Falta Pillow. Instala las dependencias con: python -m pip install -r requirements.txt") from error

        destination_folder.mkdir(parents=True, exist_ok=True)
        originals_folder.mkdir(parents=True, exist_ok=True)
        original_destination = originals_folder / self.unique_filename(originals_folder, source_path.name)
        shutil.copy2(source_path, original_destination)
        destination_name = self.unique_filename(destination_folder, f"{source_path.stem}.webp")
        destination = destination_folder / destination_name
        try:
            with Image.open(source_path) as image:
                if getattr(image, "is_animated", False):
                    raise CatalogError("No se pueden optimizar GIF animados. Convierte la imagen a JPG o PNG antes de importarla.")
                image = ImageOps.exif_transpose(image)
                image.load()
                original_width, original_height = image.size
                image.thumbnail((1600, 1600), Image.Resampling.LANCZOS)
                optimized = image.convert("RGBA") if image.mode in ("RGBA", "LA") or "transparency" in image.info else image.convert("RGB")
                optimized.save(destination, "WEBP", quality=82, method=6)
        except UnidentifiedImageError as error:
            raise CatalogError(f"El archivo no es una imagen válida: {source_path}") from error
        except OSError as error:
            raise CatalogError(f"No se pudo optimizar la imagen: {error}") from error

        self.last_image_optimization = {
            "original_bytes": source_path.stat().st_size,
            "optimized_bytes": destination.stat().st_size,
            "original_width": original_width,
            "original_height": original_height,
            "optimized_width": image.width,
            "optimized_height": image.height,
            "original_name": source_path.name,
            "optimized_name": destination.name,
        }
        return f"{public_prefix}/{destination_name}"

    @staticmethod
    def unique_filename(folder: Path, filename: str) -> str:
        candidate = Path(filename).name
        if not (folder / candidate).exists():
            return candidate
        stem, suffix = Path(candidate).stem, Path(candidate).suffix
        counter = 2
        while (folder / f"{stem}-{counter}{suffix}").exists():
            counter += 1
        return f"{stem}-{counter}{suffix}"

    @staticmethod
    def safe_key(value: str) -> str:
        normalized = unicodedata.normalize("NFKD", value.strip().lower())
        ascii_value = normalized.encode("ascii", "ignore").decode("ascii")
        key = re.sub(r"[^a-zA-Z0-9_-]+", "-", ascii_value).strip("-")
        if not key:
            raise CatalogError("La clave no puede estar vacia")
        return key

    @staticmethod
    def custom_list_key(title: str, existing_keys: set[str] | None = None) -> str:
        key = CatalogManager.safe_key(title)
        existing = existing_keys or set()
        candidate = key
        counter = 2
        while candidate in existing:
            candidate = f"{key}-{counter}"
            counter += 1
        return candidate

    @staticmethod
    def youtube_url(value: str) -> str:
        parsed = urlparse(value.strip())
        host = parsed.hostname.lower() if parsed.hostname else ""
        video_id = ""
        if host in {"youtu.be", "www.youtu.be"}:
            video_id = parsed.path.strip("/").split("/")[0]
        elif host in {"youtube.com", "www.youtube.com", "m.youtube.com"}:
            if parsed.path == "/watch":
                video_id = parse_qs(parsed.query).get("v", [""])[0]
            elif parsed.path.startswith("/shorts/"):
                video_id = parsed.path.split("/")[2]
            elif parsed.path.startswith("/embed/"):
                video_id = parsed.path.split("/")[2]
        if not re.fullmatch(r"[A-Za-z0-9_-]{6,}" , video_id):
            raise CatalogError("El enlace no parece ser una URL valida de YouTube")
        return value.strip()

    @staticmethod
    def new_custom_list(key: str, title: str, column_count: int, row_count: int) -> dict:
        if column_count < 1 or row_count < 0:
            raise CatalogError("La lista debe tener al menos una columna y cero o mas filas")
        safe_key = CatalogManager.safe_key(key)
        columns = [{"clave": f"columna-{index}", "nombre": f"Columna {index}"} for index in range(1, column_count + 1)]
        rows = [{"id": index, "uid": CatalogManager.new_uid(), "celdas": {}, "galeria": []} for index in range(1, row_count + 1)]
        return {"clave": safe_key, "titulo": title.strip(), "columnas": columns, "filas": rows}

    @staticmethod
    def add_custom_list(data: dict, custom_list: dict) -> None:
        CatalogManager.validate_custom_list(custom_list)
        lists = data.setdefault(CUSTOM_LISTS_KEY, [])
        if any(item.get("clave") == custom_list["clave"] for item in lists):
            raise CatalogError(f"Ya existe la lista '{custom_list['clave']}'")
        lists.append(copy.deepcopy(custom_list))
